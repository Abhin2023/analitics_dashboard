"""Download and parse .xlsx files from Google Drive using service account."""
import io
import os
import time
import logging
from typing import Optional
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from google_auth_httplib2 import AuthorizedHttp
import httplib2
import openpyxl

logger = logging.getLogger(__name__)

DRIVE_SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]

HTTP_TIMEOUT = 120
MAX_RETRIES = 4
RETRY_BACKOFF_BASE = 2.0

_drive_service = None


def _get_drive_service():
    global _drive_service
    if _drive_service is not None:
        return _drive_service

    sa_path = os.environ.get(
        "GOOGLE_SERVICE_ACCOUNT_JSON",
        os.path.join(os.path.dirname(__file__), "..", "..", "bp-analytics-sync-75ca3b0a13a4.json"),
    )
    if not os.path.exists(sa_path):
        raise FileNotFoundError(f"Service account JSON not found: {sa_path}")
    creds = service_account.Credentials.from_service_account_file(sa_path, scopes=DRIVE_SCOPES)
    http = AuthorizedHttp(creds, http=httplib2.Http(timeout=HTTP_TIMEOUT))
    _drive_service = build("drive", "v3", http=http, cache_discovery=False)
    return _drive_service


def _retryable(exc: BaseException) -> bool:
    return isinstance(exc, (TimeoutError, ConnectionError, ConnectionResetError, OSError))


def download_xlsx(file_id: str) -> io.BytesIO:
    """Download an xlsx file from Google Drive by file ID with retries on transient errors."""
    drive = _get_drive_service()
    last_exc = None
    for attempt in range(MAX_RETRIES):
        try:
            request = drive.files().get_media(fileId=file_id)
            buf = io.BytesIO()
            downloader = MediaIoBaseDownload(buf, request)
            done = False
            while not done:
                status, done = downloader.next_chunk()
            buf.seek(0)
            return buf
        except (TimeoutError, ConnectionError, ConnectionResetError, OSError) as e:
            last_exc = e
            if attempt < MAX_RETRIES - 1:
                wait = RETRY_BACKOFF_BASE * (2 ** attempt)
                logger.warning(
                    "Drive download retry %d/%d for %s after error %r; waiting %.1fs",
                    attempt + 1, MAX_RETRIES, file_id, e, wait,
                )
                time.sleep(wait)
    raise last_exc


def _parse_from_buf(buf: io.BytesIO):
    wb = openpyxl.load_workbook(buf, read_only=True, data_only=True)
    return wb


def parse_daily_input(file_id: str) -> list[dict]:
    """Download xlsx and parse the 'Daily Input' tab into a list of dicts."""
    buf = download_xlsx(file_id)
    return parse_daily_input_buf(buf)


def parse_daily_input_buf(buf: io.BytesIO) -> list[dict]:
    """Parse the 'Daily Input' tab from an already-downloaded xlsx buffer."""
    wb = _parse_from_buf(buf)

    if "Daily Input" not in wb.sheetnames:
        logger.error("Daily Input tab not found in xlsx")
        return []

    ws = wb["Daily Input"]
    rows = list(ws.iter_rows(values_only=True))

    if len(rows) < 5:
        return []

    data = []
    for row in rows[4:]:
        if not row or len(row) < 10:
            continue
        store_name = row[1] if row[1] else ""
        if not store_name or not str(store_name).strip():
            continue
        store_name = str(store_name).strip()

        data.append({
            "date": str(row[0]).strip() if row[0] else "",
            "store": store_name,
            "country": str(row[2]).strip() if len(row) > 2 and row[2] else "",
            "store_type": str(row[3]).strip() if len(row) > 3 and row[3] else "",
            "daily_revenue": _num(row[4]) if len(row) > 4 else None,
            "monthly_target": _num(row[5]) if len(row) > 5 else None,
            "mtd_revenue": _num(row[6]) if len(row) > 6 else None,
            "units_sold": _int(row[7]) if len(row) > 7 else None,
            "care_plus_attached": _int(row[8]) if len(row) > 8 else None,
            "prebookings": _int(row[9]) if len(row) > 9 else None,
            "ig_videos_posted": _int(row[10]) if len(row) > 10 else None,
            "ig_views_target": _num(row[11]) if len(row) > 11 else None,
            "ig_views_achieved": _num(row[12]) if len(row) > 12 else None,
            "ig_views_achd_pct": _num(row[13]) if len(row) > 13 else None,
            "ig_followers": _int(row[14]) if len(row) > 14 else None,
            "ig_new_followers": _int(row[15]) if len(row) > 15 else None,
            "ig_likes": _int(row[16]) if len(row) > 16 else None,
            "ig_comments": _int(row[17]) if len(row) > 17 else None,
            "ig_saves": _int(row[18]) if len(row) > 18 else None,
            "ig_shares": _int(row[19]) if len(row) > 19 else None,
            "ig_reposts": _int(row[20]) if len(row) > 20 else None,
            "ig_dms_received": _int(row[21]) if len(row) > 21 else None,
            "ig_manychat_handled": _int(row[22]) if len(row) > 22 else None,
            "ig_posts_published": _int(row[23]) if len(row) > 23 else None,
            "yt_views": _int(row[24]) if len(row) > 24 else None,
            "yt_likes": _int(row[25]) if len(row) > 25 else None,
            "yt_comments": _int(row[26]) if len(row) > 26 else None,
            "tt_views": _int(row[27]) if len(row) > 27 else None,
            "tt_likes": _int(row[28]) if len(row) > 28 else None,
            "tt_followers": _int(row[29]) if len(row) > 29 else None,
            "sc_views": _int(row[30]) if len(row) > 30 else None,
            "sc_shares": _int(row[31]) if len(row) > 31 else None,
            "wa_chats_received": _int(row[32]) if len(row) > 32 else None,
            "wa_walkins_booked": _int(row[33]) if len(row) > 33 else None,
            "google_rating": _num(row[34]) if len(row) > 34 else None,
            "google_new_reviews": _int(row[35]) if len(row) > 35 else None,
            "google_review_response": str(row[36]).strip() if len(row) > 36 and row[36] else "",
        })

    wb.close()
    return data


def parse_store_dashboard(file_id: str) -> list[dict]:
    """Download xlsx and parse the 'Store Dashboard' tab for summary/status data."""
    buf = download_xlsx(file_id)
    return parse_store_dashboard_buf(buf)


def parse_store_dashboard_buf(buf: io.BytesIO) -> list[dict]:
    """Parse the 'Store Dashboard' tab from an already-downloaded xlsx buffer."""
    wb = _parse_from_buf(buf)

    if "Store Dashboard" not in wb.sheetnames:
        return []

    ws = wb["Store Dashboard"]
    rows = list(ws.iter_rows(values_only=True))

    if len(rows) < 5:
        return []

    data = []
    for row in rows[4:]:
        if not row or len(row) < 5:
            continue
        store_name = row[0] if row[0] else ""
        if not store_name or not str(store_name).strip():
            continue

        data.append({
            "store": str(store_name).strip(),
            "country": str(row[1]).strip() if len(row) > 1 and row[1] else "",
            "mtd_revenue": _num(row[2]) if len(row) > 2 else None,
            "monthly_target": _num(row[3]) if len(row) > 3 else None,
            "target_pct": _num(row[4]) if len(row) > 4 else None,
            "care_plus_pct": _num(row[5]) if len(row) > 5 else None,
            "total_views": _int(row[6]) if len(row) > 6 else None,
            "engagements": _int(row[7]) if len(row) > 7 else None,
            "eng_rate_pct": _num(row[8]) if len(row) > 8 else None,
            "prebookings": _int(row[9]) if len(row) > 9 else None,
            "dms_received": _int(row[10]) if len(row) > 10 else None,
            "wa_response_pct": _num(row[11]) if len(row) > 11 else None,
            "walkins_booked": _int(row[12]) if len(row) > 12 else None,
            "insta_followers": _int(row[13]) if len(row) > 13 else None,
            "follower_growth": _int(row[14]) if len(row) > 14 else None,
            "google_rating": _num(row[15]) if len(row) > 15 else None,
            "new_reviews": _int(row[16]) if len(row) > 16 else None,
            "sales_status": str(row[17]).strip() if len(row) > 17 and row[17] else "",
            "marketing_status": str(row[18]).strip() if len(row) > 18 and row[18] else "",
        })

    wb.close()
    return data


def _num(val) -> Optional[float]:
    if val is None:
        return None
    s = str(val).strip().replace(",", "").replace("₹", "").replace("%", "").replace("—", "")
    if not s or s.lower() in ("n/a", "—", "-", "none"):
        return None
    try:
        return float(s)
    except (ValueError, TypeError):
        return None


def _int(val) -> Optional[int]:
    n = _num(val)
    return int(n) if n is not None else None
